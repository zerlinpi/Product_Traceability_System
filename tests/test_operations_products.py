"""End-to-end tests for operations-owned products and free-form purchase orders.

Covers the behaviour requested for the operations role:
  - each operations user may add their own products and only sees their own;
  - admins see every product together with its creator;
  - purchase orders are created directly from the template fields (supplier
    optional, product linked); operations only see their own orders while admins
    see all orders with the purchaser (采购人).
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app import PURCHASE_ORDER_EXPORT_COLUMNS
from capability_helpers import (
    bootstrap_admin,
    insert_user,
    login,
    make_auth_app,
    post,
    seed_scenario,
)
from test_inbound_receipts import create_order
from test_purchase_orders import setup_case


def test_operations_products_are_private_and_visible_to_admin_with_creator(tmp_path):
    app, database_path, _fake = make_auth_app(tmp_path)
    admin, admin_csrf = bootstrap_admin(app)
    # Admin's own manufacturing product (with a BOM) via the seed helper.
    seed_scenario(admin, admin_csrf, "OPS")

    insert_user(database_path, "ops.alice", "OPERATIONS")
    insert_user(database_path, "ops.bob", "OPERATIONS")
    alice, alice_csrf = login(app, "ops.alice")
    bob, bob_csrf = login(app, "ops.bob")

    # Operations register simple products with no BOM.
    alice_product = post(alice, alice_csrf, "/api/products", {"name": "Alice 产品", "modelCode": "ALICE-001"})
    assert alice_product["componentCount"] == 0
    assert alice_product["productCode"] == "ALICE-001"
    bob_product = post(bob, bob_csrf, "/api/products", {"name": "Bob 产品"})

    # Each operations user only sees their own product.
    alice_list = alice.get("/api/products").get_json()["data"]
    assert {item["id"] for item in alice_list} == {alice_product["id"]}
    bob_list = bob.get("/api/products").get_json()["data"]
    assert {item["id"] for item in bob_list} == {bob_product["id"]}

    # Admin sees every product and the creator of each.
    admin_list = admin.get("/api/products").get_json()["data"]
    admin_ids = {item["id"] for item in admin_list}
    assert {alice_product["id"], bob_product["id"]} <= admin_ids
    by_id = {item["id"]: item for item in admin_list}
    assert by_id[alice_product["id"]]["createdBy"] == "ops.alice"
    assert by_id[bob_product["id"]]["createdBy"] == "ops.bob"


def test_admin_can_create_component_less_product(tmp_path):
    # Some supplier products are complete, indivisible finished goods: an admin
    # may register a product with no BOM (empty or omitted components).
    app, database_path, _fake = make_auth_app(tmp_path)
    admin, admin_csrf = bootstrap_admin(app)

    empty_list = post(admin, admin_csrf, "/api/products", {"name": "完整成品甲", "components": []})
    assert empty_list["componentCount"] == 0
    assert empty_list["tracePlanId"] is None

    omitted = post(admin, admin_csrf, "/api/products", {"name": "完整成品乙"})
    assert omitted["componentCount"] == 0

    # Both appear in the admin product list.
    admin_ids = {item["id"] for item in admin.get("/api/products").get_json()["data"]}
    assert {empty_list["id"], omitted["id"]} <= admin_ids

    # A component-less product has no production plan, so batch generation is
    # rejected rather than silently producing an empty batch.
    rejected = admin.post(
        "/api/production-batches",
        json={"productModelId": empty_list["id"], "quantity": 1, "prefix": "NOBOM"},
        headers={"X-CSRF-Token": admin_csrf},
    )
    assert rejected.status_code == 409


def test_free_form_purchase_orders_are_scoped_to_creator(tmp_path):
    app, database_path, _fake = make_auth_app(tmp_path)
    admin, _admin_csrf = bootstrap_admin(app)

    insert_user(database_path, "ops.alice", "OPERATIONS")
    insert_user(database_path, "ops.bob", "OPERATIONS")
    alice, alice_csrf = login(app, "ops.alice")
    bob, bob_csrf = login(app, "ops.bob")

    alice_product = post(alice, alice_csrf, "/api/products", {"name": "Alice 产品"})
    bob_product = post(bob, bob_csrf, "/api/products", {"name": "Bob 产品"})

    # Alice creates a free-form order (no supplier), linked to her product.
    alice_order = post(
        alice,
        alice_csrf,
        "/api/purchase-orders",
        {
            "productModelId": alice_product["id"],
            "fields": {"SKU": "SKU-A", "供应商": "供应商甲", "实际采购量": 50, "含税单价": "9.90"},
        },
    )
    assert alice_order["supplierId"] is None
    assert alice_order["productModelId"] == alice_product["id"]
    assert alice_order["fields"]["SKU"] == "SKU-A"

    # Alice cannot link Bob's product.
    denied = alice.post(
        "/api/purchase-orders",
        json={"productModelId": bob_product["id"], "fields": {"SKU": "X"}},
        headers={"X-CSRF-Token": alice_csrf},
    )
    assert denied.status_code == 403

    bob_order = post(
        bob,
        bob_csrf,
        "/api/purchase-orders",
        {"productModelId": bob_product["id"], "fields": {"SKU": "SKU-B"}},
    )

    # Operations only see their own orders.
    alice_orders = alice.get("/api/purchase-orders").get_json()["data"]
    assert {item["id"] for item in alice_orders} == {alice_order["id"]}
    bob_orders = bob.get("/api/purchase-orders").get_json()["data"]
    assert {item["id"] for item in bob_orders} == {bob_order["id"]}

    # Admin sees all orders together with the purchaser.
    admin_orders = admin.get("/api/purchase-orders").get_json()["data"]
    admin_by_id = {item["id"]: item for item in admin_orders}
    assert {alice_order["id"], bob_order["id"]} <= set(admin_by_id)
    assert admin_by_id[alice_order["id"]]["createdBy"] == "ops.alice"
    assert admin_by_id[bob_order["id"]]["createdBy"] == "ops.bob"

    # The order exports to the template workbook.
    export = alice.get(f"/api/purchase-orders/{alice_order['id']}/export")
    assert export.status_code == 200
    assert export.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_operations_edit_copy_delete_scoped_to_creator(tmp_path):
    app, database_path, _fake = make_auth_app(tmp_path)
    bootstrap_admin(app)
    insert_user(database_path, "ops.alice", "OPERATIONS")
    insert_user(database_path, "ops.bob", "OPERATIONS")
    alice, alice_csrf = login(app, "ops.alice")
    bob, bob_csrf = login(app, "ops.bob")

    product = post(alice, alice_csrf, "/api/products", {"name": "Alice 产品"})
    order = post(
        alice,
        alice_csrf,
        "/api/purchase-orders",
        {"productModelId": product["id"], "fields": {"SKU": "A-1", "实际采购量": 10}},
    )

    # Alice edits her own order; the document number is immutable.
    edited = alice.put(
        f"/api/purchase-orders/{order['id']}",
        json={"productModelId": product["id"], "fields": {"SKU": "A-2", "实际采购量": 20}},
        headers={"X-CSRF-Token": alice_csrf},
    )
    assert edited.status_code == 200, edited.get_json()
    data = edited.get_json()["data"]
    assert data["poNo"] == order["poNo"]
    assert data["fields"]["SKU"] == "A-2"
    assert data["fields"]["实际采购量"] == 20

    # Bob can neither edit nor delete Alice's order.
    assert bob.put(
        f"/api/purchase-orders/{order['id']}",
        json={"fields": {"SKU": "X"}},
        headers={"X-CSRF-Token": bob_csrf},
    ).status_code == 403
    assert bob.delete(
        f"/api/purchase-orders/{order['id']}", headers={"X-CSRF-Token": bob_csrf}
    ).status_code == 403

    # Alice deletes her own order.
    assert alice.delete(
        f"/api/purchase-orders/{order['id']}", headers={"X-CSRF-Token": alice_csrf}
    ).status_code == 200
    assert alice.get("/api/purchase-orders").get_json()["data"] == []


def test_pushed_purchase_order_cannot_be_edited(tmp_path):
    client, _database_path, scenario, _fake = setup_case(tmp_path)
    order = create_order(client, scenario)
    assert client.post(f"/api/purchase-orders/{order['id']}/push").status_code == 200
    locked = client.put(
        f"/api/purchase-orders/{order['id']}", json={"fields": {"SKU": "X"}}
    )
    assert locked.status_code == 409
    assert "已推送" in locked.get_json()["message"]


def test_batch_export_returns_one_workbook_with_all_orders(tmp_path):
    client, _database_path, scenario, _fake = setup_case(tmp_path)
    create_order(client, scenario, 5)
    create_order(client, scenario, 8)
    response = client.get("/api/purchase-orders/export")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sheet = load_workbook(BytesIO(response.data)).active
    headers = [cell.value for cell in sheet[1]]
    assert headers == PURCHASE_ORDER_EXPORT_COLUMNS
    assert sheet.max_row == 3  # header + 2 orders


def test_lingxing_status_is_available_to_operations(tmp_path):
    app, database_path, _fake = make_auth_app(tmp_path)
    bootstrap_admin(app)
    insert_user(database_path, "ops.alice", "OPERATIONS")
    alice, _alice_csrf = login(app, "ops.alice")
    response = alice.get("/api/lingxing/status")
    assert response.status_code == 200
    data = response.get_json()["data"]
    assert set(data) == {"configured", "writeEndpointsConfigured", "endpointsConfigured"}
    # Per-operation gating: each push is enabled by its own endpoint.
    assert set(data["endpointsConfigured"]) == {
        "purchaseOrder", "inboundReceipt", "inventorySync",
    }
